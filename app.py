import streamlit as st
import os
import tempfile
import base64
from curl_cffi import requests
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# --- CONFIGURAÇÃO DA PÁGINA ---
st.set_page_config(page_title="Relatórios Eleitorais - TSE", page_icon="🗳️", layout="centered")

st.title("🗳️ Painel de Relatórios - Goiás (2024)")
st.markdown("Selecione um município para visualizar os eleitos e gerar o arquivo Word configurado.")

# Parâmetros Fixos
ANO = "2024"
ID_ELEICAO = "2045202024"
ESTADO = "GO"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

@st.cache_data
def buscar_municipios():
    url_municipios = f"https://divulgacandcontas.tse.jus.br/divulga/rest/v1/eleicao/buscar/{ESTADO}/{ID_ELEICAO}/municipios"
    try:
        resp = requests.get(url_municipios, impersonate="chrome120", timeout=20)
        if resp.status_code == 200:
            return {m.get('nome').upper(): m.get('codigo') for m in resp.json().get("municipios", [])}
    except:
        pass
    return {}

municipios_go = buscar_municipios()

if municipios_go:
    cidade_escolhida = st.selectbox("Selecione o Município:", [""] + sorted(list(municipios_go.keys())))

    if cidade_escolhida:
        if st.button(f"🔍 Gerar Relatório de {cidade_escolhida}"):
            cod_municipio = municipios_go[cidade_escolhida]
            
            with st.spinner(f"Buscando e processando dados oficiais do TSE para {cidade_escolhida}... Isso pode levar alguns segundos."):
                cargos_para_buscar = [("11", "Prefeito"), ("13", "Vereador")]
                
                temp_dir = tempfile.TemporaryDirectory()
                pasta_fotos = temp_dir.name
                eleitos = []
                
                for cod_cargo, nome_cargo in cargos_para_buscar:
                    url_lista = f"https://divulgacandcontas.tse.jus.br/divulga/rest/v1/candidatura/listar/{ANO}/{cod_municipio}/{ID_ELEICAO}/{cod_cargo}/candidatos"
                    try:
                        resp = requests.get(url_lista, impersonate="chrome120", timeout=30)
                        if resp.status_code == 200:
                            for cand in resp.json().get("candidatos", []):
                                if "Eleito" in cand.get("descricaoTotalizacao", ""):
                                    cand_id = cand["id"]
                                    url_detalhe = f"https://divulgacandcontas.tse.jus.br/divulga/rest/v1/candidatura/buscar/{ANO}/{cod_municipio}/{ID_ELEICAO}/candidato/{cand_id}"
                                    
                                    try:
                                        detalhe_resp = requests.get(url_detalhe, impersonate="chrome120", timeout=20)
                                        detalhe = detalhe_resp.json()
                                        
                                        nome_urna = detalhe.get("nomeUrna", cand.get("nomeUrna", "Sem Nome"))
                                        nome_completo = detalhe.get("nomeCompleto", cand.get("nomeCompleto", ""))
                                        partido = detalhe.get("partido", {}).get("sigla", cand.get("partido", {}).get("sigla", ""))
                                        foto_url = detalhe.get("fotoUrl")
                                        caminho_foto = os.path.join(pasta_fotos, f"{cand_id}.jpg")
                                        
                                        if foto_url:
                                            img_resp = requests.get(foto_url, impersonate="chrome120", timeout=20)
                                            if img_resp.status_code == 200:
                                                with open(caminho_foto, "wb") as f:
                                                    f.write(img_resp.content)
                                                    
                                        eleitos.append({
                                            "nome_urna": nome_urna, "nome_completo": nome_completo,
                                            "cargo": nome_cargo, "partido": partido,
                                            "foto_local": caminho_foto if os.path.exists(caminho_foto) else None
                                        })
                                        
                                        if cod_cargo == "11":
                                            for v in detalhe.get("vices", []):
                                                nome_v_urna = v.get("nomeUrna", v.get("nm_URNA", "Sem Nome"))
                                                nome_v_completo = v.get("nomeCompleto", v.get("nm_CANDIDATO", ""))
                                                partido_v = v.get("partido", "")
                                                if isinstance(partido_v, dict): partido_v = partido_v.get("sigla", "")
                                                foto_v_url = v.get("urlFoto") or v.get("fotoUrl")
                                                caminho_foto_v = os.path.join(pasta_fotos, f"vice_{cand_id}.jpg")
                                                
                                                if foto_v_url:
                                                    try:
                                                        img_resp_v = requests.get(foto_v_url, impersonate="chrome120", timeout=20)
                                                        if img_resp_v.status_code == 200:
                                                            with open(caminho_foto_v, "wb") as f:
                                                                f.write(img_resp_v.content)
                                                    except: pass
                                                eleitos.append({
                                                    "nome_urna": nome_v_urna, "nome_completo": nome_v_completo,
                                                    "cargo": "Vice-Prefeito", "partido": partido_v,
                                                    "foto_local": caminho_foto_v if os.path.exists(caminho_foto_v) else None
                                                })
                                    except: pass
                    except: pass

                if not eleitos:
                    st.error("Não foi possível encontrar eleitos para esta cidade. Tente novamente.")
                else:
                    doc = Document()
                    title = doc.add_heading(f'Eleitos de {cidade_escolhida.title()} (Gestão 2025-2028)', level=1)
                    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

                    table = doc.add_table(rows=1, cols=3)
                    table.style = 'Table Grid'
                    table.alignment = WD_ALIGN_PARAGRAPH.CENTER

                    hdr_cells = table.rows[0].cells
                    cabecalhos = ['Foto', 'Nome e Cargo', 'Partido']
                    
                    for i in range(3):
                        hdr_cells[i].text = cabecalhos[i]
                        set_cell_background(hdr_cells[i], 'D9D9D9')
                        for paragraph in hdr_cells[i].paragraphs:
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for run in paragraph.runs: run.font.bold = True

                    for v in eleitos:
                        row_cells = table.add_row().cells
                        cell_foto = row_cells[0]
                        if v.get("foto_local") and os.path.exists(v.get("foto_local")):
                            paragraph = cell_foto.paragraphs[0]
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            paragraph.add_run().add_picture(v["foto_local"], width=Inches(1.2))
                        else:
                            cell_foto.text = "Sem foto"
                            cell_foto.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

                        cell_nome = row_cells[1]
                        p_nome = cell_nome.paragraphs[0]
                        p_nome.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        run_urna = p_nome.add_run(f'"{v["nome_urna"]}"\n')
                        run_urna.font.bold = True
                        p_nome.add_run(f'{v["nome_completo"]}\n')
                        run_cargo = p_nome.add_run(v['cargo'])
                        run_cargo.font.size = Pt(9)

                        cell_partido = row_cells[2]
                        cell_partido.text = v['partido']
                        cell_partido.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

                    nome_arquivo = f'Eleitos_{cidade_escolhida.replace(" ", "_")}_GO.docx'
                    caminho_final = os.path.join(pasta_fotos, nome_arquivo)
                    doc.save(caminho_final)
                    
                    st.success("✅ Relatório finalizado com sucesso!")
                    
                    with open(caminho_final, "rb") as file:
                        st.download_button(
                            label=f"⬇️ Fazer Download de: {nome_arquivo}",
                            data=file,
                            file_name=nome_arquivo,
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                        )
else:
    st.error("Falha ao se conectar com o TSE para buscar as cidades.")
